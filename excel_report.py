"""Append a classification run to the same Excel workbook while keeping previous worksheets."""
from datetime import datetime
import math
import os
from pathlib import Path
import re
import tempfile
import unicodedata

from filelock import FileLock
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# Row 1 contains model names and row 2 contains headers; leave one blank column between groups A and B.
HEADERS = ["ID", "Requirement",
           "A label", "A reason", "A time (s)", "A error",
           "",
           "B label", "B reason", "B test focus", "B time (s)", "B error",
           "Run started"]
WIDTHS = [10, 50, 10, 50, 12, 30, 3, 10, 50, 50, 12, 30, 20]
A_COL, B_COL = 3, 8
FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)


def _set_text(cell, text):
    # External text must not be executed as an Excel formula; the complete text remains in JSON.
    cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)[:32767]
    cell.data_type = "s"


def export_excel(report, path):
    """Add each run_id only once; re-exporting does not overwrite teammates' edits to existing worksheets."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet_name = report["run_id"]
    if not re.fullmatch(r"\d{8}-\d{6}-\d{6}", sheet_name):
        raise ValueError("Invalid run ID")
    with FileLock(str(path) + ".lock", timeout=3):
        book = load_workbook(path) if path.exists() else Workbook()
        temp_path = None
        try:
            if sheet_name in book.sheetnames:
                return sheet_name
            if not path.exists():
                book.remove(book.active)
            sheet = book.create_sheet(sheet_name)
            sheet.freeze_panes = "C3"
            for col, name in ((A_COL, "A"), (B_COL, "B")):
                version = f" (prompt {report['prompt_version']})" if report.get("prompt_version") else ""
                _set_text(sheet.cell(1, col), f"{name} = {report['models'][name]}{version}")
                sheet.cell(1, col).font = BOLD
            for col, (header, width) in enumerate(zip(HEADERS, WIDTHS), 1):
                sheet.cell(2, col, header or None).font = BOLD
                sheet.column_dimensions[get_column_letter(col)].width = width
            started = datetime.fromisoformat(report["started_at"]).replace(tzinfo=None, microsecond=0)
            for rownum, item in enumerate(report["rows"], 3):
                a, b = item.get("A", {}), item.get("B", {})
                aa, ba = a.get("answer") or {}, b.get("answer") or {}
                default = "SKIPPED" if not item["requirement"].strip() else "NOT_RUN"
                values = [item["record_id"], item["requirement"],
                          a.get("label", default), aa.get("reason"),
                          a.get("elapsed_seconds"), a.get("error"),
                          None,
                          b.get("label", default), ba.get("reason"),
                          ba.get("test_focus"), b.get("elapsed_seconds"), b.get("error"),
                          started]
                for col, value in enumerate(values, 1):
                    cell = sheet.cell(rownum, col)
                    if isinstance(value, str):
                        _set_text(cell, value)
                    else:
                        cell.value = value
                    cell.font = FONT
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if col in (5, 11):
                        cell.number_format = "0.00"
                    elif col == 13:
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                lines = 1
                for value, width in zip(values, WIDTHS):
                    if isinstance(value, str):
                        lines = max(lines, sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line) / (width - 2))) for line in value.split("\n")))
                sheet.row_dimensions[rownum].height = min(409, max(16, lines * 16))
            book.active = book.index(sheet)
            fd, temp_path = tempfile.mkstemp(prefix=".report-", suffix=".xlsx", dir=path.parent)
            os.close(fd)
            book.save(temp_path)
            os.replace(temp_path, path)
            return sheet_name
        finally:
            book.close()
            if temp_path and os.path.exists(temp_path):
                os.unlink(temp_path)
