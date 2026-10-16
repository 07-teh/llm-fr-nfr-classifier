"""
Merge multiple runs on the same dataset (initial run + --skip-done resumptions) into one Task 1 results workbook without calling the API.

- For each requirement, take the first valid A and B answer (FR/NFR) in argument order and record which run it came from
- Leave the original run snapshots, databases, and team_results.xlsx unchanged
- Use B (pro + thinking mode) for the AI-only decision by default; flag A/B disagreements for human review
- Leave the human columns (Human final label, Accepted/Modified/Rejected, reason, review time) blank for teammates to fill in
- survey_label is the worksheet the respondent filled in (FR / NFR); it represents the student's original intent, not an expert answer

Usage:
    .venv/bin/python merge_runs.py results/answers/run_initial.json results/answers/run_resume.json
"""

import argparse
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
DATA_FILE = "data/healthy_nation.csv"
EXPECTED_FILE = "data/healthy_nation_expected.csv"
OUTPUT_FILE = "FINAL/task1_results.xlsx"
REVIEW_FILE = "data/healthy_nation_review.csv"  # Review subset (id, requirement), for testing with --prompts v2
AI_ONLY_FROM = "B"            # classifier used for the AI-only decision
VALID = ("FR", "NFR")
HUMAN_COLUMNS = ["Human final label", "Human decision", "Human reason", "Review time (s)"]
OLD_HUMAN_COLUMNS = ["Human+AI final", "Validation", "Justification"]  # Legacy headers, also checked before overwriting
DECISION_HELP = ("Accepted: AI label and reason are both correct; set Human final label to the AI-only label\n"
                 "Modified: label is correct but the reason is incorrect or incomplete; keep Human final label the same as AI-only and add a Human reason\n"
                 "Rejected: label is wrong; set Human final label to the opposite label and add a Human reason")

COLUMNS = [  # (header, width)
    ("ID", 9), ("Requirement", 60), ("Survey sheet", 10), ("Survey category", 22),
    ("A label", 8), ("A reason", 45), ("B label", 8), ("B test focus", 45), ("B reason", 45),
    ("A = B", 7), ("AI-only decision", 10), ("Needs review", 9),
    ("Human final label", 11), ("Human decision", 12), ("Human reason", 40), ("Review time (s)", 10),
    ("A run", 24), ("B run", 24), ("A time (s)", 9), ("B time (s)", 9),
]
V2_COLUMNS = [("v2 A label", 9), ("v2 A reason", 45), ("v2 B label", 9), ("v2 B reason", 45)]
FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)


def load_runs(run_files):
    runs = []
    for path in run_files:
        with open(SCRIPT_DIR / path, encoding="utf-8") as f:
            runs.append(json.load(f))
    return runs


def pick(runs, record_id, text, name):
    """Take the first valid answer in run order; if none exists, return the last result (ERROR, etc.). Ignore results with different requirement text."""
    last = None
    for run in runs:
        row = run["_index"].get(record_id, {})
        entry = row.get(name) if row.get("requirement") == text else None
        if entry:
            if entry.get("label") in VALID:
                return entry, run["run_id"]
            last = (entry, run["run_id"])
    return last or ({}, "")


def run_seconds(run):
    if not run.get("ended_at"):
        return None
    return (datetime.fromisoformat(run["ended_at"]) - datetime.fromisoformat(run["started_at"])).total_seconds()


def write_sheet(book, title, header, rows, widths, freeze=None):
    sheet = book.create_sheet(title)
    for col, (name, width) in enumerate(zip(header, widths), 1):
        sheet.cell(1, col, name).font = BOLD
        sheet.column_dimensions[get_column_letter(col)].width = width
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            cell = sheet.cell(r, c)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # External text is not a formula
            cell.font = FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if freeze:
        sheet.freeze_panes = freeze
    return sheet


def human_work_in(path):
    """Count populated human-review cells in an existing results workbook; do not overwrite if any contain data"""
    if not path.exists():
        return 0
    old = pd.read_excel(path, sheet_name="Task1", dtype=str, keep_default_na=False)
    return int(sum((old[c].str.strip() != "").sum() for c in HUMAN_COLUMNS + OLD_HUMAN_COLUMNS if c in old))


def main(run_files, v2_file=None, force=False):
    filled = human_work_in(SCRIPT_DIR / OUTPUT_FILE)
    if filled and not force:
        raise SystemExit(f"{OUTPUT_FILE} has {filled} populated human-review cells; merging again would overwrite them. "
                         "Save a copy or change OUTPUT_FILE first; add --force if you intend to overwrite.")
    data = pd.read_csv(SCRIPT_DIR / DATA_FILE, dtype=str, keep_default_na=False)
    expected = pd.read_csv(SCRIPT_DIR / EXPECTED_FILE, dtype=str, keep_default_na=False).set_index("id")
    runs = load_runs(run_files)
    versions = {run.get("prompt_version", "v1") for run in runs}
    if len(versions) > 1:
        raise SystemExit(f"Cannot merge runs with different prompt versions: {sorted(versions)}")
    for run in runs:
        run["_index"] = {row["record_id"]: row for row in run["rows"]}

    v2 = {}
    if v2_file:
        with open(SCRIPT_DIR / v2_file, encoding="utf-8") as f:
            v2_run = json.load(f)
        if v2_run.get("prompt_version") != "v2":
            raise SystemExit(f"{v2_file} is not a prompt v2 run")
        v2 = {r["record_id"]: r for r in v2_run["rows"]}
    columns = COLUMNS + (V2_COLUMNS if v2 else [])

    rows, stats = [], []
    for _, item in data.iterrows():
        rid, text = item["id"], item["requirement"].strip()
        exp = expected.loc[rid]
        a, a_run = pick(runs, rid, text, "A")
        b, b_run = pick(runs, rid, text, "B")
        aa, ba = a.get("answer") or {}, b.get("answer") or {}
        default = "SKIPPED" if not text else "NOT_RUN"
        a_label, b_label = a.get("label", default), b.get("label", default)
        both = a_label in VALID and b_label in VALID
        ai_only = (a_label, b_label)[AI_ONLY_FROM == "B"] if text else "SKIPPED"
        rows.append([
            rid, text, exp["expected_label"], exp["category"],
            a_label, aa.get("reason"), b_label, ba.get("test_focus"), ba.get("reason"),
            ("Yes" if a_label == b_label else "No") if both else None,
            ai_only, ("Yes" if not both or a_label != b_label else "No") if text else None,
            None, None, None, None,
            a_run, b_run, a.get("elapsed_seconds"), b.get("elapsed_seconds"),
        ])
        if v2:
            r2 = v2.get(rid) if v2.get(rid, {}).get("requirement") == text else None
            rows[-1] += ([r2["A"]["label"], (r2["A"].get("answer") or {}).get("reason"),
                          r2["B"]["label"], (r2["B"].get("answer") or {}).get("reason")]
                         if r2 and "A" in r2 and "B" in r2 else [None] * 4)
        stats.append(dict(survey=exp["expected_label"], a=a_label, b=b_label, text=bool(text)))

    s = pd.DataFrame(stats)
    s = s[s.text]
    ok = s[s.a.isin(VALID) & s.b.isin(VALID)]
    pct = lambda n, d: f"{n / d * 100:.1f}%" if d else "-"
    summary = [
        ["Non-empty requirements", len(s)],
        ["Both A and B have valid answers", len(ok)],
        ["Incomplete (ERROR / INVALID / NOT_RUN)", len(s) - len(ok)],
        ["A and B agree", f"{(ok.a == ok.b).sum()} ({pct((ok.a == ok.b).sum(), len(ok))})"],
        ["A and B disagree (human review required)", int((ok.a != ok.b).sum())],
        ["A labels FR / NFR", f"{(ok.a == 'FR').sum()} / {(ok.a == 'NFR').sum()}"],
        ["B labels FR / NFR", f"{(ok.b == 'FR').sum()} / {(ok.b == 'NFR').sum()}"],
        ["A matches survey sheet (not an expert answer)", pct((ok.a == ok.survey).sum(), len(ok))],
        ["B matches survey sheet (not an expert answer)", pct((ok.b == ok.survey).sum(), len(ok))],
        ["AI-only decision from", f"{AI_ONLY_FROM} (when A and B disagree, use {AI_ONLY_FROM}; flag for review)"],
        [],
        ["Run", "Started", "Ended", "Status", "Elapsed (s)", "A total item time (s)", "B total item time (s)"],
    ]
    for run in runs:
        sums = [round(sum(r.get(n, {}).get("elapsed_seconds") or 0 for r in run["rows"]), 1) for n in ("A", "B")]
        summary.append([run["run_id"], run["started_at"], run.get("ended_at"), run.get("status"),
                        run_seconds(run), *sums])
    summary.append([])
    summary.append(["Note", "The script processes multiple items concurrently, so total item time may exceed elapsed time; report both and explain."])

    book = Workbook()
    book.remove(book.active)
    if v2:
        summary.insert(10, ["v2 reference columns", f"{v2_file}: Only review-required items were run; this is for human review and does not change the AI-only decision"])
    summary.insert(10, ["Human review", DECISION_HELP])
    headers = [c for c, _ in columns]
    task = write_sheet(book, "Task1", headers, rows, [w for _, w in columns], freeze="C2")
    last = len(rows) + 1
    for header, options, title, help_text in (
        ("Human final label", "FR,NFR", "Final label", "Final label after human review (Human+AI decision); choose FR or NFR only"),
        ("Human decision", "Accepted,Modified,Rejected", "Human decision", DECISION_HELP),
    ):
        col = get_column_letter(headers.index(header) + 1)
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid option", error=f"Choose one of: {options}",
                            promptTitle=title, prompt=help_text[:255])
        dv.add(f"{col}2:{col}{last}")
        task.add_data_validation(dv)
    sheet = write_sheet(book, "Summary", [], [], [])
    for r, values in enumerate(summary, 1):
        for c, value in enumerate(values, 1):
            sheet.cell(r, c, value).font = FONT
    sheet.column_dimensions["A"].width = 40
    for col in "BCDEFG":
        sheet.column_dimensions[col].width = 26
    review = [(r[0], r[1]) for r in rows if r[11] == "Yes"]
    pd.DataFrame(review, columns=["id", "requirement"]).to_csv(SCRIPT_DIR / REVIEW_FILE, index=False, encoding="utf-8")
    print(f"Review required for {len(review)} items → {SCRIPT_DIR / REVIEW_FILE}")
    out = SCRIPT_DIR / OUTPUT_FILE
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    print(f"Wrote {out}")
    for line in summary[:10]:
        print("  " + ": ".join(str(v) for v in line))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Merge multiple runs into a Task 1 results workbook")
    parser.add_argument("runs", nargs="+", help="Run snapshots in priority order: results/answers/run_*.json")
    parser.add_argument("--v2", help="Run snapshot for the prompt v2 review subset; add as reference columns")
    parser.add_argument("--force", action="store_true", help="Overwrite even when human-review columns contain data")
    args = parser.parse_args()
    main(args.runs, args.v2, args.force)
