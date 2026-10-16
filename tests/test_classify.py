"""Offline regression tests: no API calls and no real credentials are used."""
import contextlib
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch, MagicMock
import pandas as pd
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import excel_report
import classify as app

CFG = dict(model=app.MODEL_A, prompt=app.PROMPT_A, thinking=False,
           temperature=0, max_tokens=None,
           required_keys=['id', 'label', 'reason'])

def answer(record_id='001'):
    return dict(id=record_id, label='FR', reason='The user records activity.')

class RegressionTests(unittest.TestCase):
    def test_csv_xlsx_keep_ids_and_literal_na(self):
        with tempfile.TemporaryDirectory() as tmp:
            df = pd.DataFrame({'id':['001'], 'requirement':['NA']})
            for suffix in ('csv', 'xlsx'):
                p = Path(tmp) / ('data.' + suffix)
                if suffix == 'csv': df.to_csv(p, index=False)
                else: df.to_excel(p, index=False)
                loaded = app.load_data(p)
                self.assertEqual(loaded.iloc[0].tolist(), ['001', 'NA'])

    def test_empty_multi_column_requirement(self):
        with patch.object(app, 'TEXT_COLUMNS', ['a', 'b']):
            self.assertEqual(app.build_requirement_text(pd.Series({'a':' ', 'b':''})), '')

    def test_schema_error_retried(self):
        invalid = answer('wrong')
        with patch.object(app, 'call_model', side_effect=[(json.dumps(invalid), None), (json.dumps(answer()), None)]) as call:
            label, entry, _, error, _ = app.classify(None, CFG, 'demo', '001')
        self.assertEqual(label, 'FR')
        self.assertIsNone(error)
        self.assertEqual(call.call_count, 2)

    def test_bad_schema_exhausts_retries(self):
        with patch.object(app, 'call_model', return_value=('{"label":"FR"}', None)):
            label, entry, _, error, _ = app.classify(None, CFG, 'demo', '001')
        self.assertEqual(label, 'INVALID')
        self.assertIsNotNone(error)

    def test_main_persists_both_independent_answers(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,Record daily steps\n')
            def response(client, model, system_prompt, user_content, thinking, **kwargs):
                data = answer()
                if 'Step 1' in system_prompt: data['test_focus'] = 'Check that steps are recorded.'
                self.assertNotIn('expected_label', user_content)
                return json.dumps(data), None
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'call_model', side_effect=response), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main(), 0)
            for name in ('a', 'b'):
                with sqlite3.connect(p/getattr(app, f'{name.upper()}_DB_FILE')) as conn:
                    self.assertEqual(conn.execute('SELECT record_id,label FROM results').fetchall(), [('001','FR')])
                entries = json.loads(next((p/app.JSON_DIR).glob(f'{name}_*.json')).read_text())
                self.assertEqual(entries[0]['record_id'], '001')

    def test_concurrent_rows_keep_order_and_skip_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n' + ''.join(f'00{n},Req {n}\n' for n in range(1, 6)) + '006,\n')
            def response(client, model, system_prompt, user_content, thinking, **kwargs):
                rid = user_content.split('\n')[0].split(': ')[1]
                data = answer(rid); data['test_focus'] = 'Check it.'
                return json.dumps(data), None
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'WORKERS', 3), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'call_model', side_effect=response), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main(), 0)
            report = json.loads(next((p/app.JSON_DIR).glob('run_*.json')).read_text())
            self.assertEqual([r['record_id'] for r in report['rows']], ['001','002','003','004','005','006'])
            self.assertTrue(all(r['A']['record_id'] == r['record_id'] == r['B']['record_id'] for r in report['rows'][:5]))
            self.assertNotIn('A', report['rows'][5])
            with sqlite3.connect(p/app.B_DB_FILE) as conn:
                self.assertEqual(conn.execute('select count(*) from results').fetchone()[0], 5)

    def test_fatal_error_still_keeps_finished_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,One\n002,Two\n003,Three\n')
            real = app.classify
            def fake(client, cfg, user_content, record_id):
                if record_id == '002': raise RuntimeError('DeepSeek HTTP 402')
                with patch.object(app, 'call_model', return_value=(json.dumps(dict(answer(record_id), test_focus='Check it.')), None)):
                    return real(client, cfg, user_content, record_id)
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'WORKERS', 3), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'classify', side_effect=fake), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main(), 1)
            report = json.loads(next((p/app.JSON_DIR).glob('run_*.json')).read_text())
            self.assertTrue(report['status'].startswith('stopped: '))
            self.assertEqual([r['record_id'] for r in report['rows'] if 'B' in r], ['001', '003'])

    def test_fatal_error_with_queued_rows_does_not_hang(self):
        import threading
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,One\n002,Two\n003,Three\n004,Four\n')
            real = app.classify
            def fake(client, cfg, user_content, record_id):
                if record_id == '002': raise RuntimeError('DeepSeek HTTP 402')
                with patch.object(app, 'call_model', return_value=(json.dumps(dict(answer(record_id), test_focus='Check it.')), None)):
                    return real(client, cfg, user_content, record_id)
            out = {}
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'WORKERS', 1), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'classify', side_effect=fake), contextlib.redirect_stdout(io.StringIO()):
                t = threading.Thread(target=lambda: out.setdefault('code', app.main()), daemon=True)
                t.start(); t.join(10)
            self.assertFalse(t.is_alive(), 'Hung after a fatal error')
            self.assertEqual(out['code'], 1)
            report = json.loads(next((p/app.JSON_DIR).glob('run_*.json')).read_text())
            got = [r['record_id'] for r in report['rows'] if 'B' in r]
            self.assertIn('001', got); self.assertNotIn('002', got)

    def test_skip_done_reruns_only_unfinished_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,One\n002,Two\n003,Three\n')
            (p/app.JSON_DIR).mkdir(parents=True)
            ok = lambda rid, text='One': dict(record_id=rid, requirement=text, A=dict(label='FR'), B=dict(label='NFR'))
            (p/app.JSON_DIR/'run_old.json').write_text(json.dumps(dict(rows=[ok('001'), dict(record_id='002', requirement='Two', A=dict(label='FR'), B=dict(label='ERROR')), ok('003', 'Old text')])))
            seen = []
            def response(client, model, system_prompt, user_content, thinking, **kwargs):
                rid = user_content.split('\n')[0].split(': ')[1]; seen.append(rid)
                return json.dumps(dict(answer(rid), test_focus='Check it.')), None
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'call_model', side_effect=response), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main([app.JSON_DIR + '/run_old.json']), 0)
            self.assertEqual(sorted(set(seen)), ['002', '003'])
            report = json.loads(next(f for f in (p/app.JSON_DIR).glob('run_*.json') if f.name != 'run_old.json').read_text())
            self.assertEqual([r['record_id'] for r in report['rows']], ['002', '003'])
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'PROMPT_VERSION', 'v2'), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaisesRegex(SystemExit, 'cannot resume with prompt'): app.main([app.JSON_DIR + '/run_old.json'])

    def test_report_records_prompt_version_and_v2_loads(self):
        import prompts_v2
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,One\n')
            prompts = []
            def response(client, model, system_prompt, user_content, thinking, **kwargs):
                prompts.append(system_prompt)
                return json.dumps(dict(answer(), test_focus='Check it.')), None
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'COMMON_PROMPT', prompts_v2.COMMON_PROMPT), patch.object(app, 'PROMPT_A', prompts_v2.PROMPT_A), patch.object(app, 'PROMPT_B', prompts_v2.PROMPT_B), patch.object(app, 'PROMPT_VERSION', prompts_v2.VERSION), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'call_model', side_effect=response), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main(), 0)
            report = json.loads(next((p/app.JSON_DIR).glob('run_*.json')).read_text())
            self.assertEqual(report['prompt_version'], 'v2')
            self.assertIn('R10.', report['prompts']['common'])
            self.assertTrue(all('R10.' in sp for sp in prompts))

    def test_duplicate_id_stops_before_api(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,One\n001,Two\n')
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI') as client, contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaisesRegex(SystemExit, 'duplicated'): app.main()
                client.assert_not_called()



class ExcelTests(unittest.TestCase):
    def report(self, run_id='20260930-100000-000001'):
        return dict(run_id=run_id, started_at='2026-09-30T10:00:00+08:00',
                    status='completed', models={'A':'model-a', 'B':'model-b'},
                    rows=[dict(record_id='001', requirement='=1+1',
                               A=dict(label='FR', answer=answer(), elapsed_seconds=1.25,
                                      time='2026-09-30 10:00:01'),
                               B=dict(label='ERROR', answer=None, error='Response timed out'))])

    def test_append_keeps_previous_sheet_and_reexport_is_idempotent(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'team.xlsx'
            first = self.report()
            app.export_excel(first, path)
            book = load_workbook(path)
            book.active['P3'] = 'Teammate note'
            book.save(path)
            book.close()
            app.export_excel(self.report('20260930-100100-000002'), path)
            app.export_excel(first, path)
            book = load_workbook(path)
            self.assertEqual(len(book.sheetnames), 2)
            self.assertEqual(book.worksheets[0]['P3'].value, 'Teammate note')
            sheet = book.worksheets[1]
            self.assertEqual([c.value or '' for c in sheet[2]], excel_report.HEADERS)
            self.assertEqual(sheet['A3'].value, '001')
            self.assertEqual(sheet['B3'].value, '=1+1')
            self.assertEqual(sheet['B3'].data_type, 's')
            self.assertEqual(sheet['L3'].value, 'Response timed out')
            self.assertEqual(sheet['E3'].value, 1.25)
            self.assertEqual((sheet['C1'].value, sheet['H1'].value), ('A = model-a', 'B = model-b'))
            self.assertEqual(str(sheet['M3'].value), '2026-09-30 10:00:00')
            self.assertEqual(sheet.freeze_panes, 'C3')
            self.assertEqual(len(sheet.tables), 0)
            self.assertEqual(len(sheet.merged_cells.ranges), 0)
            self.assertIsNone(sheet['G3'].value)
            self.assertEqual(sheet['B3'].font.name, 'Calibri')
            self.assertTrue(sheet["A2"].font.bold)
            self.assertIsNone(sheet["A2"].fill.patternType)
            book.close()

    def test_failed_atomic_save_keeps_original(self):
        import excel_report
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'team.xlsx'
            app.export_excel(self.report(), path)
            original=path.read_bytes()
            with patch.object(excel_report.os, 'replace', side_effect=PermissionError):
                with self.assertRaises(PermissionError):
                    app.export_excel(self.report('20260930-100100-000002'), path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(tmp).glob('.report-*')), [])

    def test_excel_failure_keeps_json_sqlite_and_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)
            (p/'data.csv').write_text('id,requirement\n001,Record steps\n')
            def response(client, model, system_prompt, user_content, thinking, **kwargs):
                data=answer()
                data['test_focus']='Check recording.'
                return json.dumps(data), None
            with patch.object(app, 'SCRIPT_DIR', p), patch.object(app, 'DATA_FILE', 'data.csv'), patch.object(app, 'load_dotenv'), patch.dict(app.os.environ, {'DEEPSEEK_API_KEY':'test-only'}), patch.object(app, 'OpenAI'), patch.object(app, 'call_model', side_effect=response), patch.object(app, 'export_excel', side_effect=PermissionError), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(app.main(), 1)
            report=json.loads(next((p/app.JSON_DIR).glob('run_*.json')).read_text())
            self.assertEqual(report['rows'][0]['A']['label'], 'FR')
            self.assertEqual(report['rows'][0]['B']['label'], 'FR')
            with sqlite3.connect(p/app.A_DB_FILE) as conn:
                self.assertEqual(conn.execute('select count(*) from results').fetchone()[0], 1)
            app.export_excel(report,p/'team_results.xlsx')
            self.assertTrue((p/'team_results.xlsx').exists())

if __name__ == '__main__': unittest.main()
